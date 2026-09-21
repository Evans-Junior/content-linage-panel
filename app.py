from dotenv import load_dotenv

import os
load_dotenv(os.environ.get(".env"))

import base64
import hmac
import io
import json
import re

import anthropic
import streamlit as st
from google import genai
from google.genai import types
from pptx import Presentation

CLAUDE_MODEL = "claude-sonnet-5"
# Override with GEMINI_MODEL in .env if a newer Flash model is available to you.
GEMINI_MODEL = os.environ.get("GEMINI_MODEL") or "gemini-3.6-flash"  # `or` also covers an empty GEMINI_MODEL=

REQUIRED_VARS = ["ANTHROPIC_API_KEY", "GEMINI_API_KEY", "APP_PASSWORD"]

st.set_page_config(page_title="Content Lineage Panel")

# ---- Startup checks and password gate --------------------------------------

missing = [v for v in REQUIRED_VARS if not os.environ.get(v)]
if missing:
    st.error(
        "Missing required environment variable(s): " + ", ".join(missing)
        + "."
    )
    st.stop()

st.title("Content Lineage Panel")

if not st.session_state.get("authed"):
    entered = st.text_input("Password", type="password")
    if entered:
        if hmac.compare_digest(entered.encode(), os.environ["APP_PASSWORD"].encode()):
            st.session_state["authed"] = True
            st.rerun()
        else:
            st.error("Incorrect password.")
    st.stop()

claude = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
gemini = genai.Client(api_key=os.environ["GEMINI_API_KEY"])


# ---- Reading the uploaded files ---------------------------------------------

def slides_to_text(upload) -> str:
    """Plain-text version of the slide deck (PPTX or text outline)."""
    if upload.name.lower().endswith(".pptx"):
        parts = []
        for i, slide in enumerate(Presentation(io.BytesIO(upload.getvalue())).slides, 1):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    lines.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        lines.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide:
                lines.append("Notes: " + slide.notes_slide.notes_text_frame.text)
            parts.append(f"--- Slide {i} ---\n" + "\n".join(l for l in lines if l.strip()))
        return "\n\n".join(parts)
    return upload.getvalue().decode("utf-8", errors="replace")


# ---- Check 1: text points (Claude, with citations) --------------------------

TEXT_PROMPT = """Below is the text of a slide deck that was generated from the attached source document.

<slides>
{slides}
</slides>

Identify points from the source document that are MISSING from the slides (dropped) or that
appear only in a much shorter, weaker or less specific form (condensed). Only report points of
real substance, not trivia. Report at most 12, most important first.

Use exactly this format for every finding, with no other text before, between or after them:

@@@ <short title>
STATUS: dropped OR condensed - <one line saying what was lost>
SLIDE: <for condensed points, the slide number(s) from the "--- Slide N ---" labels where the shortened version appears; NONE if dropped or the deck has no slide labels>
ORIGINAL: <quote the relevant source passage>

Quote the source passage so it is cited. If nothing important is missing, reply with only: NONE"""


def text_check(name: str, data: bytes, slides: str) -> list[dict]:
    doc = {
        "type": "document",
        "source": {
            "type": "base64",
            "media_type": "application/pdf",
            "data": base64.standard_b64encode(data).decode(),
        },
        "title": name,
        "citations": {"enabled": True},
    }

    response = claude.messages.create(
        model=CLAUDE_MODEL,
        max_tokens=16000,
        messages=[{
            "role": "user",
            "content": [doc, {"type": "text", "text": TEXT_PROMPT.format(slides=slides)}],
        }],
    )

    # The model's text arrives in blocks; cited blocks carry a `citations` list.
    # Split on the @@@ marker, attaching each block's citations to the finding it belongs to.
    findings = []
    for block in response.content:
        if block.type != "text":
            continue
        pieces = block.text.split("@@@")
        for n, piece in enumerate(pieces):
            if n > 0:
                findings.append({"text": "", "citations": []})
            if not findings:
                continue
            findings[-1]["text"] += piece
            if n == len(pieces) - 1 and block.citations:
                findings[-1]["citations"].extend(block.citations)

    results = []
    for f in findings:
        title = f["text"].strip().splitlines()[0].strip() if f["text"].strip() else "Untitled"
        status = re.search(r"STATUS:\s*(dropped|condensed)\s*[-–—:]?\s*(.*)", f["text"], re.I)
        slide = re.search(r"SLIDE:\s*(.*)", f["text"], re.I)
        slide = slide.group(1).strip() if slide else ""
        # Page and wording come from the API's citations, not from the model's own text.
        pages, quotes = [], []
        for c in f["citations"]:
            quotes.append(c.cited_text)
            if c.type == "page_location":
                start, end = c.start_page_number, c.end_page_number - 1
                pages.append(str(start) if start >= end else f"{start}-{end}")
        results.append({
            "title": title,
            "location": ", ".join(dict.fromkeys(pages)) or "not cited",
            "status": status.group(1).lower() if status else "unknown",
            "note": status.group(2).strip() if status else "",
            "slide": "" if slide.upper().startswith("NONE") else slide,
            "quotes": quotes,
            "source": "Text check",
        })
    return results


# ---- Check 2: figures (Gemini reads PDF pages as images) --------------------

FIGURE_PROMPT = """You are given a source PDF and the text of a slide deck generated from it.

<slides>
{slides}
</slides>

List every figure, chart or diagram in the PDF. For each, decide whether its content (the data or
the main takeaway) made it into the slides. Return a JSON array where each item has:
  "title": short title of the figure,
  "page": the PDF page number (1 = first page),
  "status": one of "present", "condensed", "dropped",
  "note": one line explaining the status,
  "slide": the slide number(s) from the "--- Slide N ---" labels where it appears, or "" if none,
  "description": what the figure shows, including key numbers or labels.
Return [] if the PDF has no figures."""


def figure_check(pdf_data: bytes, slides: str) -> list[dict]:
    response = gemini.models.generate_content(
        model=GEMINI_MODEL,
        contents=[
            types.Part.from_bytes(data=pdf_data, mime_type="application/pdf"),
            FIGURE_PROMPT.format(slides=slides),
        ],
        config=types.GenerateContentConfig(response_mime_type="application/json"),
    )
    results = []
    for item in json.loads(response.text):
        if item.get("status") == "present":
            continue
        results.append({
            "title": item.get("title", "Untitled figure"),
            "location": str(item.get("page", "?")),
            "status": item.get("status", "unknown"),
            "note": item.get("note", ""),
            "slide": str(item.get("slide", "")),
            "quotes": [item.get("description", "")],
            "source": "Figure check (less precise)",
        })
    return results


# ---- UI ---------------------------------------------------------------------

source = st.file_uploader(
    "Source document (PDF only)", type=["pdf"],
    help="Word users: File > Save as PDF first. Page numbers in the results count from the "
         "first page of the PDF file, which may differ from any numbers printed on the pages.",
)
slides_upload = st.file_uploader(
    "Generated slide deck", type=["pptx", "txt", "md"],
    help="PPTX slides are numbered automatically. A text outline has no slide numbers "
         "unless you label them yourself, e.g. 'Slide 1', 'Slide 2'.",
)

if source and slides_upload and st.button("Check lineage"):
    pdf_data = source.getvalue()
    if not pdf_data.startswith(b"%PDF"):
        st.error("That file is not a valid PDF. Please export your document as a PDF and upload it again.")
        st.stop()
    slides = slides_to_text(slides_upload)

    results, errors = [], []
    with st.spinner("Running text check..."):
        try:
            results += text_check(source.name, pdf_data, slides)
        except Exception as e:
            errors.append(f"Text check failed: {e}")
    with st.spinner("Running figure check..."):
        try:
            results += figure_check(pdf_data, slides)
        except Exception as e:
            errors.append(f"Figure check failed: {e}")

    st.session_state["results"] = results
    st.session_state["errors"] = errors

for err in st.session_state.get("errors", []):
    st.error(err)

if "results" in st.session_state:
    results = st.session_state["results"]
    st.subheader(f"{len(results)} omitted or condensed item(s)")
    for r in results:
        with st.container(border=True):
            st.markdown(f"**{r['title']}** · PDF page {r['location']}")
            slide_note = f" (now on slide {r['slide']})" if r["slide"] else ""
            st.markdown(f"**{r['status'].capitalize()}**: {r['note']}{slide_note}")
            st.caption(f"Source: {r['source']}")
            with st.expander("Original wording" if r["source"] == "Text check" else "Figure description"):
                for q in r["quotes"] or ["(no citation returned)"]:
                    st.markdown(f"> {q}")
